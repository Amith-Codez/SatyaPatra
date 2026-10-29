"""Builds the SatyaPatra SIH 2026 idea deck on the official SIH2026 template.

Keeps the template's logos, footer, slide titles and the "idea details pointers"
word for word (template rule 5), removes the instructions slide (template note),
and lays our content out under each pointer.

Run:  python build_deck.py   ->  SatyaPatra_SIH26261_Idea.pptx
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

TEAM_NAME = "[Team Name]"   # fill before export: must not contain the college name
TEAM_ID = "[Team ID]"

NAVY, BLUE, INK, MUTED = "1F3864", "0070C0", "1B2333", "5B6476"
LINE, BG, WHITE = "D5DCE8", "F4F7FB", "FFFFFF"
OK, BAD, WARN = "1A7F4B", "C62828", "A85F00"
FONT = "Arial"


def rgb(h):
    return RGBColor.from_string(h)


def remove_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def set_runs_text(paragraph, text):
    """Replace a paragraph's text but keep the first run's formatting."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run().text = text
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def bullet(p, level=0):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.16 + 0.16 * level))))
    pPr.set("indent", str(int(-Inches(0.14))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
    pPr.append(bu)


def no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("a:buNone")) is None:
        pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def text(slide, x, y, w, h, paras, size=11, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=3, name=None):
    """paras: list of items. Each item is a str, or dict with keys
    runs=[(text, {bold,color,size,italic,link,underline})], bullet=bool, align, size, color, bold."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    first = True
    for item in paras:
        if isinstance(item, str):
            item = {"runs": [(item, {})]}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = item.get("align", align)
        p.space_after = Pt(item.get("space_after", space_after))
        if item.get("bullet"):
            bullet(p, item.get("level", 0))
        else:
            no_bullet(p)
        for t, o in item["runs"]:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", item.get("size", size)))
            f.bold = o.get("bold", item.get("bold", bold))
            f.italic = o.get("italic", False)
            f.underline = o.get("underline", False)
            f.color.rgb = rgb(o.get("color", item.get("color", color)))
            if o.get("link"):
                r.hyperlink.address = o["link"]
    return tb


def box(slide, x, y, w, h, fill=BG, line=LINE, radius=0.06, line_w=1.0, name=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    if name:
        shp.name = name
    return shp


def circle(slide, x, y, d, fill, label, size=11):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = rgb(fill)
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = rgb(WHITE)
    return c


def pointer(slide, x, y, w, label, size=12.5, color=BLUE, h=0.3):
    """Template 'idea details pointer', kept word for word, styled as a section header."""
    return text(slide, x, y, w, h, [{"runs": [(label, {"bold": True, "color": color, "size": size})]}])


def table(slide, x, y, w, col_w, rows, header_fill=NAVY, size=9.5, row_h=0.3, first_col_bold=True,
          col_colors=None):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    for attr in ("bandRow", "firstRow"):
        tblPr.set(attr, "0")
    # drop the default table style so our fills win
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        tblPr.remove(style)
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows):
        tbl.rows[ri].height = Inches(row_h)
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if ri == 0:
                cell.fill.fore_color.rgb = rgb(header_fill)
            else:
                cell.fill.fore_color.rgb = rgb(WHITE if ri % 2 else BG)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            if isinstance(val, tuple):
                val, vcolor = val
            else:
                vcolor = None
            r = p.add_run()
            r.text = val
            r.font.name = FONT
            r.font.size = Pt(size)
            if ri == 0:
                r.font.bold = True
                r.font.color.rgb = rgb(WHITE)
            else:
                r.font.bold = bool(first_col_bold and ci == 0) or bool(vcolor)
                c = vcolor or (col_colors[ci] if col_colors else INK)
                r.font.color.rgb = rgb(c)
    return gt


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def delete_slide(prs, index):
    sldIdLst = prs.slides._sldIdLst
    sldId = sldIdLst[index]
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)


def shape_by_name(slide, prefix):
    for sh in slide.shapes:
        if sh.name.startswith(prefix):
            return sh
    raise KeyError(prefix)


def set_team_oval(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().replace("\n", " ") in ("Your Team Name",):
            p = sh.text_frame.paragraphs[0]
            set_runs_text(p, TEAM_NAME)
            for r in p.runs:
                r.font.size = Pt(12)
            for extra in sh.text_frame.paragraphs[1:]:
                extra._p.getparent().remove(extra._p)


def set_title(slide, s):
    t = slide.shapes.title
    paras = t.text_frame.paragraphs
    set_runs_text(paras[0], s)
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)


